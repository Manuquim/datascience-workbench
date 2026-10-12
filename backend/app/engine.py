import csv
import hashlib
import json
import math
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

import duckdb
from openpyxl import load_workbook


def identifier(name):
    if not isinstance(name, str) or not name or len(name) > 200 or "\x00" in name:
        raise ValueError("Nombre de columna inválido.")
    return '"' + name.replace('"', '""') + '"'


def literal(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError("Los valores numéricos deben ser finitos.")
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (bytes, bytearray)):
        return value.hex()
    if isinstance(value, (list, tuple)):
        return [json_value(x) for x in value]
    if isinstance(value, dict):
        return {str(k): json_value(v) for k, v in value.items()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def rows_json(cursor):
    columns = [d[0] for d in cursor.description]
    return [dict(zip(columns, [json_value(v) for v in row])) for row in cursor.fetchall()]


def file_hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(4 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class Engine:
    def __init__(self, settings, store, project_id, workspace):
        self.settings, self.store, self.project_id = settings, store, project_id
        self.workspace = Path(workspace)
        self.workspace.mkdir(parents=True, exist_ok=True)
        self.con = duckdb.connect(
            config={
                "memory_limit": settings.memory_limit,
                "threads": settings.threads,
                "temp_directory": str(self.workspace / "spill"),
                "max_temp_directory_size": settings.temp_limit,
                "autoload_known_extensions": False,
                "autoinstall_known_extensions": False,
            }
        )

    def close(self):
        self.con.close()

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()

    def path(self, data):
        path = (self.store.root / data["path"]).resolve()
        if not path.is_relative_to(self.store.root.resolve()) or not path.is_file():
            raise ValueError("Archivo no disponible dentro del almacenamiento del proyecto.")
        return path

    def source(self, data):
        path = literal(str(self.path(data)))
        if data["format"] == "parquet":
            return f"read_parquet({path})"
        if data["format"] == "csv":
            return f"read_csv({path}, auto_detect=true, header=true, sample_size=20480)"
        if data["format"] in ("json", "jsonl", "ndjson"):
            return f"read_json_auto({path})"
        raise ValueError("Formato de datos no soportado.")

    def view(self, data, name="input_data"):
        self.con.execute(
            f"CREATE OR REPLACE TEMP VIEW {identifier(name)} AS SELECT * FROM {self.source(data)}"
        )
        return identifier(name)

    def schema(self, relation):
        return [
            {"name": r[0], "type": r[1]}
            for r in self.con.execute(f"DESCRIBE SELECT * FROM {relation}").fetchall()
        ]

    def preview(self, relation, limit=50, offset=0):
        return rows_json(
            self.con.execute(f"SELECT * FROM {relation} LIMIT {int(limit)} OFFSET {int(offset)}")
        )

    def profile(self, relation, metadata=None, quality=None):
        from .statistics import profile

        return profile(self, relation, metadata, quality)

    def compile_step(self, relation, step):
        schema = self.schema(relation)
        names = [c["name"] for c in schema]

        def col(name):
            if name not in names:
                raise ValueError(f"No existe la columna: {name}")
            return identifier(name)

        def new_alias(name):
            identifier(name)
            if name.casefold() in {n.casefold() for n in names}:
                raise ValueError(f"La columna ya existe: {name}")
            return identifier(name)

        def projection(replacements):
            return ", ".join(replacements.get(n, col(n)) for n in names)

        op = step.op
        warning = None
        if op == "select":
            if not step.columns or len(set(step.columns)) != len(step.columns):
                raise ValueError("Selecciona columnas diferentes.")
            sql = f"SELECT {', '.join(col(c) for c in step.columns)} FROM {relation}"
        elif op == "filter":
            c = col(step.column)
            operator = step.operator
            if operator in ("is_null", "not_null"):
                condition = f"{c} IS {'NOT ' if operator == 'not_null' else ''}NULL"
            elif operator == "contains":
                if step.value is None:
                    raise ValueError("Indica el texto a buscar.")
                condition = f"contains(CAST({c} AS VARCHAR), {literal(step.value)})"
            else:
                if operator not in {"eq", "ne", "gt", "ge", "lt", "le"} or step.value is None:
                    raise ValueError("Filtro incompleto; usa is_null para los valores ausentes.")
                symbol = {"eq": "=", "ne": "<>", "gt": ">", "ge": ">=", "lt": "<", "le": "<="}[
                    operator
                ]
                condition = f"{c} {symbol} {literal(step.value)}"
            sql = f"SELECT * FROM {relation} WHERE {condition}"
        elif op == "deduplicate":
            if step.columns:
                raise ValueError("La eliminación de duplicados inicial compara todas las columnas.")
            sql = f"SELECT DISTINCT * FROM {relation}"
        elif op == "drop_nulls":
            if not step.columns:
                raise ValueError("Indica las columnas en las que quitar nulos.")
            sql = f"SELECT * FROM {relation} WHERE " + " AND ".join(
                f"{col(c)} IS NOT NULL" for c in step.columns
            )
        elif op == "fill_null":
            c = col(step.column)
            if step.value is None:
                raise ValueError("Indica un valor de sustitución.")
            sql = f"SELECT {projection({step.column: f'COALESCE({c}, {literal(step.value)}) AS {c}'})} FROM {relation}"
        elif op == "cast":
            c = col(step.column)
            if not step.target_type:
                raise ValueError("Indica el tipo de destino.")
            sql = f"SELECT {projection({step.column: f'CAST({c} AS {step.target_type}) AS {c}'})} FROM {relation}"
        elif op == "rename":
            c, alias = col(step.column), new_alias(step.alias)
            sql = f"SELECT {projection({step.column: f'{c} AS {alias}'})} FROM {relation}"
        elif op in ("join", "concat"):
            other = self.store.dataset(self.project_id, step.dataset_id)
            right = self.view(other, "right_data")
            right_names = [c["name"] for c in self.schema(right)]
            if op == "concat":
                if names != right_names:
                    raise ValueError(
                        "Concatenar requiere las mismas columnas y orden; normalízalas primero."
                    )
                sql = f"SELECT * FROM {relation} UNION ALL SELECT * FROM {right}"
            else:
                left_key = col(step.left_key)
                if step.right_key not in right_names:
                    raise ValueError("La clave de la tabla derecha no existe.")
                right_key = identifier(step.right_key)
                duplicates = self.con.execute(
                    f"SELECT EXISTS(SELECT 1 FROM {right} WHERE {right_key} IS NOT NULL "
                    f"GROUP BY {right_key} HAVING COUNT(*) > 1)"
                ).fetchone()[0]
                if duplicates and not step.allow_many_to_many:
                    raise ValueError(
                        "La clave derecha tiene duplicados: la unión puede multiplicar filas. "
                        "Revísala; habilita allow_many_to_many solo si esa multiplicación es intencionada."
                    )
                if duplicates:
                    warning = (
                        "Unión con claves derechas duplicadas autorizada; puede multiplicar filas."
                    )
                output = [f"l.{identifier(n)}" for n in names]
                used = {n.casefold() for n in names}
                for n in right_names:
                    if n == step.right_key:
                        continue
                    alias = n if n.casefold() not in used else n + "_right"
                    while alias.casefold() in used:
                        alias += "_right"
                    used.add(alias.casefold())
                    output.append(f"r.{identifier(n)} AS {identifier(alias)}")
                sql = (
                    f"SELECT {', '.join(output)} FROM {relation} l {step.how.upper()} JOIN {right} r "
                    f"ON l.{left_key}=r.{right_key}"
                )
        elif op == "aggregate":
            groups = [col(c) for c in step.columns]
            if not step.metrics:
                raise ValueError("Indica al menos una métrica.")
            aliases = [m.alias.casefold() for m in step.metrics] + [
                c.casefold() for c in step.columns
            ]
            if len(set(aliases)) != len(aliases):
                raise ValueError("Las métricas y agrupaciones deben tener nombres diferentes.")
            metrics = []
            for m in step.metrics:
                value = "*" if m.function == "count" and not m.column else col(m.column)
                metrics.append(f"{m.function.upper()}({value}) AS {identifier(m.alias)}")
            sql = f"SELECT {', '.join(groups + metrics)} FROM {relation}"
            if groups:
                sql += " GROUP BY " + ", ".join(groups)
        elif op == "date_part":
            c, alias = col(step.column), new_alias(step.alias)
            if not step.part:
                raise ValueError("Indica la parte de fecha.")
            sql = f"SELECT *, date_part({literal(step.part)}, CAST({c} AS TIMESTAMP)) AS {alias} FROM {relation}"
        elif op == "ratio":
            c, d, alias = col(step.column), col(step.denominator), new_alias(step.alias)
            sql = f"SELECT *, CAST({c} AS DOUBLE) / NULLIF(CAST({d} AS DOUBLE), 0) AS {alias} FROM {relation}"
            warning = "Los cocientes con denominador cero o nulo producen NULL."
        elif op == "sort":
            if not step.columns:
                raise ValueError("Indica columnas de ordenación.")
            sql = f"SELECT * FROM {relation} ORDER BY " + ", ".join(
                f"{col(c)} {'DESC' if step.descending else 'ASC'} NULLS LAST" for c in step.columns
            )
        else:
            raise ValueError("Operación no permitida.")
        return sql, warning

    def execute_plan(self, data, plan, progress):
        from .metadata import derive, dictionary

        if plan.questions:
            raise ValueError("Resuelve las preguntas del plan antes de ejecutarlo.")
        relation = self.view(data)
        metadata = dictionary(data)
        for name, meta in metadata.items():
            meta.setdefault("lineage", []).append(
                {
                    "operation": "source",
                    "dataset_id": data["id"],
                    "dictionary_revision": data.get("dictionary_revision", 0),
                    "column": name,
                }
            )
        history = []
        for i, step in enumerate(plan.steps):
            before = self.con.execute(f"SELECT count(*) FROM {relation}").fetchone()[0]
            progress(i, f"Paso {i + 1}/{len(plan.steps)}: {step.op}")
            sql, warning = self.compile_step(relation, step)
            # Each stage is materialized to disk to avoid huge in-memory Python DataFrames.
            path = self.workspace / f"step-{i:02d}.parquet"
            self.con.execute(
                f"COPY ({sql}) TO {literal(str(path))} (FORMAT PARQUET, COMPRESSION ZSTD)"
            )
            name = f"step_{i}"
            self.con.execute(
                f"CREATE TEMP VIEW {identifier(name)} AS SELECT * FROM read_parquet({literal(str(path))})"
            )
            relation = identifier(name)
            other = (
                dictionary(self.store.dataset(self.project_id, step.dataset_id))
                if step.dataset_id
                else None
            )
            metadata = derive(metadata, step, self.schema(relation), other)
            count = self.con.execute(f"SELECT COUNT(*) FROM {relation}").fetchone()[0]
            history.append(
                {
                    "step": step.model_dump(),
                    "rows": count,
                    "rows_before": before,
                    "rows_delta": count - before,
                    "warning": warning,
                }
            )
        progress(len(plan.steps), "Calculando calidad y estadísticas finales")
        output = self.workspace / "result.parquet"
        self.con.execute(
            f"COPY (SELECT * FROM {relation}) TO {literal(str(output))} (FORMAT PARQUET, COMPRESSION ZSTD)"
        )
        profile = self.profile(relation, metadata)
        for chart in plan.charts:
            self.chart(relation, chart, limit=1000, metadata=metadata)
        return output, profile, history, metadata

    def chart(self, relation, spec, limit=5000, metadata=None, dataset=None):
        from .statistics import calculate_chart

        return calculate_chart(self, relation, spec, limit, metadata, dataset)


def excel_to_csv(source, destination, sheet=None):
    # openpyxl reads worksheet rows lazily but shared strings/styles can be loaded eagerly.
    # Bound expanded archive size and these eager metadata parts before opening a workbook.
    with ZipFile(source) as archive:
        infos = archive.infolist()
        if sum(item.file_size for item in infos) > 512 * 1024**2:
            raise ValueError(
                "Excel: contenido descomprimido superior a 512 MB; utiliza CSV o Parquet."
            )
        for item in infos:
            limit = 64 * 1024**2 if item.filename.endswith("sharedStrings.xml") else 8 * 1024**2
            if (
                item.filename.endswith(("sharedStrings.xml", "styles.xml"))
                and item.file_size > limit
            ):
                raise ValueError("Excel: metadatos demasiado grandes; utiliza CSV o Parquet.")
            if item.flag_bits & 1:
                raise ValueError("Excel cifrado no admitido.")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        if worksheet.max_column and worksheet.max_column > 200:
            raise ValueError("Excel: máximo 200 columnas en esta versión.")
        with destination.open("w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index > 1_048_576:
                    raise ValueError("Excel: demasiadas filas.")
                writer.writerow(
                    [v.isoformat() if isinstance(v, (date, datetime)) else v for v in row]
                )
        return worksheet.title
    finally:
        workbook.close()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
